import io

from docx import Document

from backend.app.store import Store
from .conftest import finish


def test_api_runs_when_frontend_build_is_incomplete(tmp_path,monkeypatch):
    from backend.app import main
    from fastapi.testclient import TestClient
    monkeypatch.delenv('CONTEXTGUARD_API_ONLY',raising=False)
    monkeypatch.setattr(main,'ROOT',tmp_path)
    dist=tmp_path/'frontend'/'dist';dist.mkdir(parents=True)
    (dist/'index.html').write_text('<html></html>')
    with TestClient(main.create_app(str(tmp_path/'partial.sqlite3'))) as api_client:
        assert api_client.get('/api/health').status_code==200
        assert api_client.get('/').status_code==404


def test_sample_report_and_brief(client, sample):
    assert len(sample["sources"]) == 4
    assert len(sample["report"]["findings"]) == 4
    assert sample["report"]["rejected_count"] == 0
    assert all(x["provider"] == "fixture" for x in sample["report"]["traces"])
    brief = client.get(f"/api/cases/{sample['id']}/brief")
    assert brief.status_code == 200
    assert "Fictional fixture replay" in brief.text
    assert "insufficient evidence" in brief.text


def test_resolution_persistence_recheck_and_undo(client, sample):
    prefix = f"/api/cases/{sample['id']}"
    finding = next(f for f in sample["report"]["findings"] if f["check_id"] == "dates")
    response = client.post(prefix + "/resolutions", json={"version": sample["version"], "finding_id": finding["id"], "action": "confirm", "answer": "November 20 is the revised release date, confirmed by the client.", "owner": "Maya"})
    assert response.status_code == 202
    updated = finish(client, sample["id"])
    f = next(x for x in updated["report"]["findings"] if x["id"] == finding["id"])
    assert f["status"] == "user_confirmed" and f["owner"] == "Maya"
    assert sum(x["status"] == "open" for x in updated["report"]["findings"]) == 3
    assert updated["changes"]
    reopened = Store(client.app.state.store.path).get(sample["id"])
    assert reopened.resolutions[0].answer.startswith("November 20")
    response = client.post(prefix + f"/resolutions/{updated['resolutions'][0]['id']}/undo", json={"version": updated["version"]})
    assert response.status_code == 202
    undone = finish(client, sample["id"])
    assert all(x["status"] == "open" for x in undone["report"]["findings"])
    assert undone["resolutions"][0]["revoked"]


def test_unknown_answer_does_not_resolve(client, sample):
    response = client.post(f"/api/cases/{sample['id']}/resolutions", json={"version": sample["version"], "finding_id": sample["report"]["findings"][0]["id"], "action": "unknown", "answer": "", "owner": "Client"})
    assert response.status_code == 202
    case = finish(client, sample["id"])
    assert all(x["status"] == "open" for x in case["report"]["findings"])


def test_stale_revision_rejected(client, sample):
    response = client.post(f"/api/cases/{sample['id']}/analyze", json={"version": 0})
    assert response.status_code == 409


def test_sample_inputs_cannot_be_changed(client, sample):
    response = client.post(f"/api/cases/{sample['id']}/sources/text", json={"version": sample["version"], "name": "Extra", "text": "Conflicting extra data"})
    assert response.status_code == 400


def test_live_sources_failure_and_duplicate_handling(client):
    case = client.post("/api/cases", json={"title": "My review", "goal": "Check our implementation kickoff"}).json()
    prefix = f"/api/cases/{case['id']}"
    assert client.post(prefix + "/analyze", json={"version": case["version"]}).status_code == 400
    response = client.post(prefix + "/sources/text", json={"version": case["version"], "name": "Requirements", "text": "The final release is November 15.\n\nOwner: Maya."})
    assert response.status_code == 200
    case = response.json()
    assert case["sources"][0]["passages"][1]["text"] == "Owner: Maya."
    assert client.post(prefix + "/sources/text", json={"version": case["version"], "name": "Duplicate", "text": "The final release is November 15.\n\nOwner: Maya."}).status_code == 400
    assert client.post(prefix + "/analyze", json={"version": case["version"]}).status_code == 400
    response = client.post(prefix + f"/sources/upload?version={case['version']}", files={"file": ("broken.pdf", b"not a pdf", "application/pdf")})
    assert response.status_code == 200
    case = response.json()
    assert case["sources"][-1]["status"] == "failed"
    assert not case["sources"][-1]["passages"]
    response = client.request("DELETE", prefix + "/sources/" + case["sources"][0]["id"], json={"version": case["version"]})
    assert response.status_code == 200 and len(response.json()["sources"]) == 1


def test_docx_upload_preserves_table(client):
    doc = Document(); doc.add_paragraph("Approved scope."); doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Deadline: November 15"
    data = io.BytesIO(); doc.save(data)
    case = client.post("/api/cases", json={}).json()
    response = client.post(f"/api/cases/{case['id']}/sources/upload?version={case['version']}", files={"file": ("scope.docx", data.getvalue())})
    assert response.status_code == 200
    assert any("Deadline: November 15" in p["text"] for p in response.json()["sources"][0]["passages"])


def test_criteria_exclusion_has_no_false_rejection(client, sample):
    criteria = sample["criteria"]
    criteria[0]["excluded"] = True
    criteria[0]["exclusion_reason"] = "Release planning is outside the kickoff scope."
    response = client.put(f"/api/cases/{sample['id']}/criteria", json={"version": sample["version"], "criteria": criteria})
    assert response.status_code == 200 and response.json()["report"] is None
    case = response.json()
    client.post(f"/api/cases/{case['id']}/analyze", json={"version": case["version"]})
    case = finish(client, case["id"])
    assert case["report"]["rejected_count"] == 0
    assert not any(f["check_id"] == "dates" for f in case["report"]["findings"])


def test_cross_origin_and_host_rejected(client):
    assert client.post("/api/cases", json={}, headers={"Origin": "https://evil.test"}).status_code == 403
    assert client.get("/api/health", headers={"Host": "evil.test"}).status_code == 403


def test_delete_case_and_missing_case(client, sample):
    assert client.request("DELETE", f"/api/cases/{sample['id']}", json={"version": sample["version"]}).status_code == 200
    assert client.get(f"/api/cases/{sample['id']}").status_code == 404


def test_blank_resolution_rejected(client, sample):
    response = client.post(f"/api/cases/{sample['id']}/resolutions", json={"version": sample["version"], "finding_id": sample["report"]["findings"][0]["id"], "action": "dismiss", "answer": " "})
    assert response.status_code == 400


def test_brief_rejects_failed_recheck(client, sample):
    case = client.app.state.store.get(sample["id"])
    case.run.status = "failed"
    client.app.state.store.save(case)
    assert client.get(f"/api/cases/{case.id}/brief").status_code == 409
