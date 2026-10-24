"""Independent source readers. No adapter owns project routing or project memory."""
import io
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import httpx

from .contracts import Passage


class AdapterFailure(Exception):
    def __init__(self, message, retry_after=None):
        super().__init__(message)
        self.retry_after = retry_after


def manual(filename, content):
    filename = Path(filename.replace('\\', '/')).name[:150]
    if not filename or len(content) > 4 * 1024 * 1024:
        raise AdapterFailure('Select a named file smaller than 4 MB, or paste the relevant text.')
    suffix = Path(filename).suffix.lower()
    author, timestamp = None, None
    try:
        if suffix in {'.txt', '.md'}:
            texts = [(f'{filename}, passage {i + 1}', text.strip())
                     for i, text in enumerate(content.decode('utf-8-sig').split('\n\n')) if text.strip()]
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            if reader.metadata:
                author = reader.metadata.author or None
                timestamp = reader.metadata.modification_date or reader.metadata.creation_date
            if len(reader.pages) > 30:
                raise AdapterFailure('PDF exceeds 30 pages. Select an excerpt or paste relevant text.')
            texts = [(f'{filename}, page {i + 1}', page.extract_text() or '') for i, page in enumerate(reader.pages)]
        elif suffix == '.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if len(archive.infolist()) > 1000 or sum(f.file_size for f in archive.infolist()) > 20 * 1024 * 1024:
                    raise AdapterFailure('DOCX expanded content is too large. Paste the relevant text.')
            doc = Document(io.BytesIO(content))
            author = doc.core_properties.author or None
            timestamp = doc.core_properties.modified or doc.core_properties.created
            strings = [p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
            texts = [(f'{filename}, paragraph/table row {i + 1}', text) for i, text in enumerate(strings) if text.strip()]
        else:
            raise AdapterFailure('Supported formats: TXT, MD, PDF and DOCX. Paste text as a fallback.')
        if len(texts) > 300 or sum(len(t) for _, t in texts) > 60000 or any(len(t) > 30000 for _, t in texts):
            raise AdapterFailure('Document text exceeds review limits. Select a smaller excerpt.')
        if not texts:
            raise AdapterFailure('No extractable text. Scanned PDFs need a text export; paste text as a fallback.')
        return [Passage(item_id=filename + ':' + str(i + 1), text=t, location=loc, author=author, timestamp=timestamp,
                        details={'document_name': filename, 'timestamp_basis': 'document metadata' if timestamp else 'unavailable'},
                        extraction='complete' if t.strip() else 'failed',
                        uncertainty='' if t.strip() else 'No text extracted from this page.')
                for i, (loc, t) in enumerate(texts)]
    except AdapterFailure:
        raise
    except Exception:
        raise AdapterFailure('Document parsing failed. Check the file or paste the text instead.') from None


class Vexa:
    def __init__(self, base, key, transport=None):
        self.client = httpx.Client(base_url=base.rstrip('/') + '/', headers={'X-API-Key': key},
                                   timeout=20, trust_env=False, transport=transport)

    @staticmethod
    def validate(platform, meeting_id):
        patterns = {'google_meet': r'[a-z]{3}-[a-z]{4}-[a-z]{3}', 'teams': r'\d{8,20}', 'zoom': r'\d{9,11}'}
        if platform not in patterns or not re.fullmatch(patterns[platform], meeting_id):
            raise AdapterFailure('Choose a supported platform and its native meeting ID.')

    def request(self, method, path, **kwargs):
        try:
            response = self.client.request(method, path, **kwargs)
            response.raise_for_status()
            return response.json()
        except (httpx.HTTPError, ValueError):
            raise AdapterFailure('Vexa request failed. Check access, meeting admission and service status.') from None

    def join(self, platform, meeting_id, consent, passcode=None):
        if consent is not True:
            raise AdapterFailure('Explicit meeting authorization and participant consent are required.')
        self.validate(platform, meeting_id)
        payload = {'platform': platform, 'native_meeting_id': meeting_id, 'bot_name': 'ContextGuard — authorized notes', 'recording_enabled': False}
        if passcode:
            payload['passcode'] = passcode
        self.request('POST', 'bots', json=payload)
        return {'status': 'requested', 'note': 'Joining is not confirmed. A host may need to admit the bot.'}

    def stop(self, platform, meeting_id):
        self.validate(platform, meeting_id)
        self.request('DELETE', f'bots/{platform}/{meeting_id}')
        return {'status': 'stop_requested', 'note': 'The Vexa stop request was accepted.'}

    def fetch(self, platform, meeting_id):
        self.validate(platform, meeting_id)
        payload = self.request('GET', f'transcripts/{platform}/{meeting_id}')
        rows = payload.get('segments', [])
        if not rows:
            raise AdapterFailure('No transcript is available yet. Check admission or try again; no content was invented.')
        output = []
        for i, segment in enumerate(rows):
            start = segment.get('start', segment.get('start_time'))
            finish = segment.get('end', segment.get('end_time'))
            completed = segment.get('completed', False)
            output.append(Passage(item_id=str(segment.get('segment_id') or f'{start}:{i}'), text=segment.get('text') or '',
                author=segment.get('speaker') or None, external_identity=segment.get('speaker_id'),
                timestamp=segment.get('absolute_start_time') or segment.get('created_at'),
                location=f'{platform}/{meeting_id}, {start}–{finish}s', link=payload.get('constructed_meeting_url') or '',
                details={'meeting_id': meeting_id, 'platform': platform, 'start': start, 'end': finish, 'completed': completed},
                extraction='complete' if completed and payload.get('status') in {'completed', 'finished'} else 'partial',
                confidence=segment.get('confidence'), uncertainty='Speaker is unattributed.' if not segment.get('speaker') else ''))
        # Preserve the complete selected transcript, not credentials, recordings or private join passcodes.
        original = {k: payload.get(k) for k in ('id', 'platform', 'native_meeting_id', 'status', 'start_time', 'end_time', 'segments')}
        return output, original


class Slack:
    def __init__(self, token, team, transport=None):
        self.team = team
        self.client = httpx.Client(base_url='https://slack.com/api/', headers={'Authorization': 'Bearer ' + token},
                                   timeout=15, trust_env=False, transport=transport)

    def call(self, method, **params):
        try:
            response = self.client.get(method, params=params)
            if response.status_code == 429:
                raise AdapterFailure('Slack rate limit reached. Retry after the indicated delay.', response.headers.get('Retry-After'))
            response.raise_for_status()
            payload = response.json()
            if not payload.get('ok'):
                code = payload.get('error')
                safe = code if code in {'token_revoked', 'invalid_auth', 'missing_scope', 'not_in_channel', 'channel_not_found'} else 'request_failed'
                raise AdapterFailure('Slack: ' + safe + '. Check the authorized workspace and channel access.')
            return payload
        except (httpx.HTTPError, ValueError):
            raise AdapterFailure('Slack connection failed. No content was imported.') from None

    def fetch(self, channel, thread=None, start=None, end=None):
        if not re.fullmatch(r'[CG][A-Z0-9]{2,30}', channel):
            raise AdapterFailure('Provide the selected channel ID.')
        auth = self.call('auth.test')
        if auth.get('team_id') != self.team:
            raise AdapterFailure('Token workspace does not match the selected source.')
        self.call('conversations.info', channel=channel)
        params = {'channel': channel, 'limit': 100, 'inclusive': 'true'}
        if thread:
            if not re.fullmatch(r'\d+\.\d+', thread):
                raise AdapterFailure('Provide a valid selected thread timestamp.')
            params['ts'] = thread
        if start:
            params['oldest'] = start
        if end:
            params['latest'] = end
        messages, complete = [], False
        for _ in range(3):
            page = self.call('conversations.replies' if thread else 'conversations.history', **params)
            messages += page.get('messages', [])
            cursor = page.get('response_metadata', {}).get('next_cursor', '')
            if not cursor and not page.get('has_more'):
                complete = True
                break
            if not cursor:
                break
            params['cursor'] = cursor
        passages = []
        for message in messages:
            m = message.get('message', message) if message.get('subtype') == 'message_changed' else message
            ts = m.get('deleted_ts') if m.get('subtype') == 'message_deleted' else m.get('ts')
            if not ts:
                continue
            permalink = ''
            if not m.get('subtype') == 'message_deleted':
                try:
                    permalink = self.call('chat.getPermalink', channel=channel, message_ts=ts).get('permalink', '')
                except AdapterFailure:
                    pass
            passages.append(Passage(item_id=ts, text=m.get('text') or '', author=m.get('user') or m.get('bot_id'),
                external_identity=m.get('user'), timestamp=datetime.fromtimestamp(float(ts), timezone.utc),
                location=f'Slack {self.team}/{channel}/{ts}', link=permalink,
                details={'team_id': self.team, 'channel_id': channel, 'thread_ts': m.get('thread_ts'), 'edited': m.get('edited')},
                extraction='complete' if complete else 'partial', deleted=m.get('subtype') == 'message_deleted',
                uncertainty='' if complete else 'History pull is truncated; additional messages were not analyzed.'))
        return passages, {'team': self.team, 'channel': channel, 'thread': thread, 'complete': complete, 'messages': messages}


def baileys_events(events, selected_chat):
    """Test/import boundary only. No socket, pairing, send or history-sync capability."""
    passages = []
    for event in events:
        key = event.get('key', {})
        if key.get('remoteJid') != selected_chat:
            continue  # Never retain unselected chats, including in an original event.
        message = event.get('message') or {}
        text = message.get('conversation') or message.get('extendedTextMessage', {}).get('text') or ''
        passages.append(Passage(item_id=str(key.get('id') or ''), text=text,
            author=event.get('pushName') or None, external_identity=key.get('participant') or selected_chat,
            timestamp=datetime.fromtimestamp(int(event['messageTimestamp']), timezone.utc) if event.get('messageTimestamp') else None,
            location=f'Selected WhatsApp chat {selected_chat}, message {key.get("id")}',
            details={'chat_id': selected_chat, 'message_id': key.get('id')},
            extraction='complete' if text else 'failed', uncertainty='Baileys test/import boundary; live linked-device operation is disabled.'))
    if not passages:
        raise AdapterFailure('No selected-chat text events. Import a text export instead.')
    return passages
