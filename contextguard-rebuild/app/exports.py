"""Deterministic document exports; never fill gaps with invented prose."""
import io


def consolidated(service, actor, review_id):
    review = service.get_review(actor, review_id)
    from .service import now, require
    require(not review['stale'] and review['status'] == 'complete', 'Rerun stale or failed reviews before exporting.', 409)
    project = service.project(actor, review['project_id'])
    users = {u['id']: u['username'] for u in service.store.all('user', actor['workspace'])}
    lines = ['# ' + project['name'] + ' — consolidated context', '', 'Generated: ' + now(),
             'Goal: ' + review['goal'], 'Review status: ' + review['readiness'],
             'Mode: ' + str(review['provider']), '', '## Confirmed current decisions']
    decisions = service.decisions(actor, project['id'])
    active = [d for d in decisions if not d['superseded_by']]
    if not active:
        lines.append('No decisions have been confirmed. Source statements below are not approved decisions.')
    for decision in active:
        lines += ['', decision['text'], 'Confirmed by ' + users.get(decision['confirmed_by'], 'Project member') +
                  ' on ' + decision['confirmed_at'] + '. Decision: ' + decision['id'],
                  'Supporting/context evidence: ' + ', '.join(c['evidence_id'] for c in decision['citations'])]
        if not decision['support_current']:
            lines.append('Attention: supporting source evidence has changed; reconfirm this decision.')
    lines += ['', '## Unresolved questions']
    unresolved = [f for f in review['findings'] if f['state'] == 'open' and f['kind'] != 'verified_fact']
    if not unresolved:
        lines.append('No open findings in this review. This is not a guarantee of complete or correct context.')
    for finding in unresolved:
        lines += ['', '### ' + finding['title'], finding['question'],
                  'Accountable person: ' + users.get(finding['assignee_id'], 'Unassigned'),
                  'Impact: ' + finding['impact'], 'Next action: ' + finding['next_action']]
        for response in finding['responses']:
            lines.append('Unconfirmed response: ' + response['answer'])
    lines += ['', '## Original source evidence',
              'These are source statements as received, including older or conflicting statements. '
              'Confirmed decisions above state what the team has agreed; the original evidence is not rewritten.']
    for evidence in review['evidence']:
        lines += ['', '### ' + evidence['id'], evidence['location'],
                  'Author/speaker: ' + (evidence['author'] or 'Unknown') + '; date: ' + (evidence['timestamp'] or 'Unknown'),
                  evidence['text'], 'Source: ' + (evidence['link'] or 'Local import; inspect this evidence ID in ContextGuard')]
    lines += ['', '## Limitations'] + ['- ' + x for x in review['failures'] + review['limitations']]
    return '\n'.join(lines)


def docx_bytes(markdown):
    from docx import Document
    from docx.shared import Pt, RGBColor
    document = Document()
    style = document.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(11)
    style.font.color.rgb = RGBColor.from_string('243D33')
    for line in markdown.splitlines():
        if line.startswith('#'):
            level = min(len(line) - len(line.lstrip('#')), 3)
            document.add_heading(line.lstrip('# ').strip(), level=level)
        elif line:
            document.add_paragraph(line[2:] if line.startswith('- ') else line,
                                   style='List Bullet' if line.startswith('- ') else None)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()
